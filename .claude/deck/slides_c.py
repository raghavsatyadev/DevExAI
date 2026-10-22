"""Demo slides: the end-to-end video and real screenshots from a test run (saved AI answers)."""
import os

from pptx.util import Inches, Pt

from helpers import (LINE, MARGIN, MUTED, NAVY, W, blank, header, notes, step, text, _tag)

DEMO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "demo")
VIDEO = os.path.join(DEMO, "DevExAI-flow.mp4")  # not in git (too big); the video is also on Drive
TEST_RUN = ("Recorded on a test Slack workspace and a test dashboard. The AI answers were saved "
            "replays, so no live Claude calls were made. Screens are unedited except that a "
            "temporary folder path and an older test post are hidden.")


def shot(slide, name, x, y, w, caption=None):
    """Place demo/<name>.jpg at width w, keeping its shape, with a caption under it."""
    pic = slide.shapes.add_picture(os.path.join(DEMO, name + ".jpg"), Inches(x), Inches(y),
                                   width=Inches(w))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    _tag(pic)
    h = pic.height / 914400
    if caption:
        text(slide, x, y + h + 0.06, w, 0.3, caption, size=14, color=MUTED)
    return h


def s_video(prs, n):
    s = blank(prs)
    header(s, "Demo", "The whole flow in one minute", n)
    x, y, w, h = MARGIN, 1.6, 9.0, 9.0 * 9 / 16
    poster = os.path.join(DEMO, "poster.jpg")
    if os.path.exists(VIDEO):
        s.shapes.add_movie(VIDEO, Inches(x), Inches(y), Inches(w), Inches(h),
                           poster_frame_image=poster, mime_type="video/mp4")
    else:
        s.shapes.add_picture(poster, Inches(x), Inches(y), Inches(w), Inches(h))
    rx = x + w + 0.35
    step(1)
    text(s, rx, y, W - MARGIN - rx, 4.0, [
        [("Idea to installed tool", {"bold": True, "color": NAVY})],
        "Posted in Slack, checked, planned, approved, built, reviewed, published, installed.",
        [("Test run", {"bold": True, "color": NAVY})],
        "Test workspace, saved AI answers, no live Claude calls.",
        [("Google Slides", {"bold": True, "color": NAVY})],
        "Plays in PowerPoint. In Slides, open DevExAI-flow.mp4 from the DevExAI Drive folder.",
    ], size=16, spacing=10)
    notes(s, "Video of one idea going all the way through DevExAI, about one minute. " + TEST_RUN)


def s_slack(prs, n):
    s = blank(prs)
    header(s, "Demo", "What people see in Slack", n)
    step(1)
    shot(s, "thread", MARGIN, 1.5, 3.1, "The idea's thread")
    x, w, y = MARGIN + 3.6, 8.1, 1.6
    for k, (name, cap) in enumerate([
        ("idea", "1. Someone posts an idea and gets a tracking number."),
        ("approve", "2. The maintainer gets the plan in a direct message, with buttons."),
        ("announce", "3. After review and Publish, the tool is announced in #ai-tools."),
    ], start=2):
        step(k)
        y += shot(s, name, x, y, w, cap) + 0.95
    notes(s, "Real Slack screens from the test run. Left: the idea's thread with the safety "
             "check and plan. Right: the post, the maintainer's buttons, the announcement. "
             + TEST_RUN)


def s_built(prs, n):
    s = blank(prs)
    header(s, "Demo", "The built tool, its blog post, and installing it", n)
    step(1)
    shot(s, "files", MARGIN, 1.5, 4.9, "Dashboard: files written by the AI and its review")
    step(2)
    shot(s, "blog", MARGIN + 5.3, 1.5, 4.7, "Blog post with the install steps")
    step(3)
    shot(s, "terminal", MARGIN, 4.45, 6.9, "Real PowerShell output on a test computer")
    text(s, MARGIN + 7.3, 4.45, W - 2 * MARGIN - 7.3, 2.2, [
        "devexai install sets up tools made of .claude, .gemini or .codex files.",
        "This one is a single script, so it prints its own steps and is copied by hand.",
    ], size=14, color=MUTED, spacing=8)
    notes(s, "Left: the idea page on the dashboard. Right: the public blog post. Bottom: "
             "devexai list and install, then the script prints the branch. " + TEST_RUN)


SLIDES = [s_video, s_slack, s_built]
