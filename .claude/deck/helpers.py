"""Drawing helpers for the DevExAI deck. Arial; photos only where they carry meaning."""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "Arial"
W, H = 13.333, 7.5
MARGIN = 0.6

NAVY = RGBColor(0x0F, 0x17, 0x2A)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x47, 0x55, 0x69)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
LINE = RGBColor(0xCB, 0xD5, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x43, 0x38, 0xCA)
ACCENT_SOFT = RGBColor(0xE0, 0xE7, 0xFF)
GREEN = RGBColor(0x0F, 0x76, 0x6E)
GREEN_SOFT = RGBColor(0xCC, 0xFB, 0xF1)
AMBER = RGBColor(0xB4, 0x53, 0x09)
AMBER_SOFT = RGBColor(0xFE, 0xF3, 0xC7)
RED = RGBColor(0xB9, 0x1C, 0x1C)
RED_SOFT = RGBColor(0xFE, 0xE2, 0xE2)
SKY = RGBColor(0x93, 0xC5, 0xFD)

PHOTOS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "photos")

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    cp = prs.core_properties
    cp.title, cp.author, cp.last_modified_by, cp.comments = "DevExAI", "", "", ""
    return prs


_STEP = [None]


def step(k):
    """Shapes drawn after this fade in as build step k (None = static).
    anim.ps1 reads the "buildNN" name prefix."""
    _STEP[0] = k


def _tag(shape):
    if _STEP[0] is not None:
        shape.name = f"build{_STEP[0]:02d} {shape.name}"
    return shape


def blank(prs, bg=WHITE):
    step(None)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return slide


def _style_run(run, size, bold, color, italic=False, font=FONT):
    f = run.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = color


def text(slide, x, y, w, h, paras, size=16, bold=False, color=INK, align="l",
         anchor="t", spacing=4, line=1.1):
    """paras: str, or list of paragraphs; a paragraph is a str or a list of
    runs, where a run is a str or (text, {"bold":..,"color":..,"size":..})."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ANCHOR[anchor]
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        p.line_spacing = line
        if i:
            p.space_before = Pt(spacing)
        runs = [para] if isinstance(para, str) else para
        for r in runs:
            t, opt = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = t
            _style_run(run, opt.get("size", size), opt.get("bold", bold),
                       opt.get("color", color), opt.get("italic", False),
                       opt.get("font", FONT))
    return _tag(box)


def rect(slide, x, y, w, h, fill=LIGHT, line=None, radius=0.06, shape=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    return _tag(s)


def photo(slide, name, x, y, w, h):
    """Place photos/<name>.jpg, cropped (never stretched) to fill the frame."""
    pic = slide.shapes.add_picture(os.path.join(PHOTOS, name + ".jpg"), Inches(x),
                                   Inches(y), Inches(w), Inches(h))
    iw, ih = pic.image.size
    src, frame = iw / ih, w / h
    if src > frame:
        pic.crop_left = pic.crop_right = (1 - frame / src) / 2
    else:
        pic.crop_top = pic.crop_bottom = (1 - src / frame) / 2
    return _tag(pic)


def shade(slide, color=NAVY, opacity=0.8):
    """Full-slide tint so white text stays readable over a photo."""
    s = rect(slide, 0, 0, W, H, fill=color, radius=0)
    clr = s._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    alpha = clr.makeelement(qn("a:alpha"), {"val": str(int(opacity * 100000))})
    clr.append(alpha)
    return s


def header(slide, kicker, title, number=None):
    text(slide, MARGIN, 0.42, 9, 0.3, kicker.upper(), size=12, bold=True, color=ACCENT)
    text(slide, MARGIN, 0.72, W - 2 * MARGIN, 0.7, title, size=30, bold=True, color=NAVY)
    if number is not None:
        text(slide, W - MARGIN - 1, H - 0.45, 1, 0.3, str(number), size=11,
             color=MUTED, align="r")
        text(slide, MARGIN, H - 0.45, 6, 0.3, "DevExAI, proposal draft", size=11,
             color=MUTED)


def card(slide, x, y, w, h, title, body, accent=ACCENT, fill=LIGHT, title_size=18,
         body_size=14, badge=None):
    rect(slide, x, y, w, h, fill=fill)
    rect(slide, x, y + 0.18, 0.07, h - 0.36, fill=accent, radius=0)
    top = y + 0.2
    if badge:
        text(slide, x + 0.3, top, w - 0.5, 0.4, badge, size=20, bold=True, color=accent)
        top += 0.45
    text(slide, x + 0.3, top, w - 0.5, 0.45, title, size=title_size, bold=True, color=NAVY)
    text(slide, x + 0.3, top + 0.5, w - 0.5, h - (top - y) - 0.6, body, size=body_size,
         color=MUTED)


def chip(slide, x, y, w, h, label, fill=ACCENT_SOFT, color=ACCENT, size=14, bold=True):
    rect(slide, x, y, w, h, fill=fill, radius=0.3)
    text(slide, x + 0.08, y, w - 0.16, h, label, size=size, bold=bold, color=color,
         align="c", anchor="m")


def arrow(slide, x, y, w=0.35, h=0.3, color=MUTED, down=False):
    kind = MSO_SHAPE.DOWN_ARROW if down else MSO_SHAPE.RIGHT_ARROW
    rect(slide, x, y, w, h, fill=color, shape=kind)


def table(slide, x, y, widths, rows, size=14, head_size=14, row_h=0.5):
    """rows[0] is the header row. A cell is a str or (str, {"bold":..,"color":..})."""
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y),
                                   Inches(sum(widths)), Inches(row_h * len(rows)))
    tbl = shape.table
    tbl.first_row = True
    for i, wd in enumerate(widths):
        tbl.columns[i].width = Inches(wd)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            t, opt = (val, {}) if isinstance(val, str) else val
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (WHITE if r % 2 else LIGHT)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            run = p.add_run()
            run.text = t
            if r == 0:
                _style_run(run, head_size, True, WHITE)
            else:
                _style_run(run, opt.get("size", size), opt.get("bold", c == 0),
                           opt.get("color", NAVY if c == 0 else INK))
    return _tag(shape)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body
